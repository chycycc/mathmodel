# -*- coding: utf-8 -*-
import os
import openpyxl

base_dir = '数模题目/D题/数据/无人机应急物资运输基础数据'
files = os.listdir(base_dir)

with open('scratch/excel_summary.txt', 'w', encoding='utf-8') as out:
    for f in sorted(files):
        if not f.endswith('.xlsx'):
            continue
        out.write(f"\n==================== FILE: {f} ====================\n")
        fpath = os.path.join(base_dir, f)
        wb = openpyxl.load_workbook(fpath, data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            out.write(f"\n--- Sheet: {sname} (rows: {ws.max_row}, cols: {ws.max_column}) ---\n")
            for r in range(1, ws.max_row + 1):
                row_vals = [str(ws.cell(r, c).value if ws.cell(r, c).value is not None else '') for c in range(1, ws.max_column + 1)]
                if any(v.strip() for v in row_vals):
                    out.write(' | '.join(row_vals) + '\n')

print("Saved excel summary successfully!")
