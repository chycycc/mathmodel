import openpyxl, os, glob
base=r'D:/python project/mathmodel/数模题目/D题/数据/无人机应急物资运输基础数据'
for f in glob.glob(base+'/*.xlsx'):
 print('FILE',os.path.basename(f))
 wb=openpyxl.load_workbook(f,data_only=True)
 for ws in wb.worksheets:
  print(' SHEET',ws.title,ws.max_row,ws.max_column)
  for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,5),values_only=True): print('  ',row)
