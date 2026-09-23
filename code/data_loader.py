# -*- coding: utf-8 -*-
"""
数据加载与三维地理空间计算核心模块
包含：
1. 节点（O01、S001~S015）信息与经纬度/海拔
2. 运输与中继无人机性能、电池、通信射频参数
3. 80个离散物资货箱详细清单
4. 30米 DEM 栅格剖面、航段最高海拔与净空巡航高度计算
5. 航段飞行耗时与动态能耗精确计算
6. 两阶段电池等效充电模型
7. 空间三维射线追踪视距遮挡（LOS）与双向通信链路可用性判定
"""

import os
import math
import numpy as np
import scipy.io
import openpyxl

# 自适应工程工作区根路径
_CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
_WORKSPACE_ROOT = os.path.dirname(_CURRENT_DIR)
DATA_DIR = os.path.join(_WORKSPACE_ROOT, "数模题目", "D题", "数据")
BASE_DATA_DIR = os.path.join(DATA_DIR, "无人机应急物资运输基础数据")
GEO_DATA_DIR = os.path.join(DATA_DIR, "镇龙乡地理空间数据", "镇龙乡及周边地理数据")

# 1. 加载 DEM 矩阵
dem_mat_path = os.path.join(GEO_DATA_DIR, "数字高程模型数据（DEM）", "镇龙乡及周边30米DEM.mat")
mat_data = scipy.io.loadmat(dem_mat_path)
DEM_GRID = mat_data['dem']
DEM_LATS = mat_data['latitude'].flatten()
DEM_LONS = mat_data['longitude'].flatten()
LAT_STEP = DEM_LATS[1] - DEM_LATS[0]
LON_STEP = DEM_LONS[1] - DEM_LONS[0]

def get_dem_elevation(lon, lat):
    """查询指定经纬度的地面海拔（双线性插值）"""
    r = (lat - DEM_LATS[0]) / LAT_STEP
    c = (lon - DEM_LONS[0]) / LON_STEP
    r0 = int(math.floor(r))
    c0 = int(math.floor(c))
    r1 = min(r0 + 1, DEM_GRID.shape[0] - 1)
    c1 = min(c0 + 1, DEM_GRID.shape[1] - 1)
    r0 = max(0, min(r0, DEM_GRID.shape[0] - 1))
    c0 = max(0, min(c0, DEM_GRID.shape[1] - 1))
    dr = r - r0
    dc = c - c0
    e00 = DEM_GRID[r0, c0]
    e01 = DEM_GRID[r0, c1]
    e10 = DEM_GRID[r1, c0]
    e11 = DEM_GRID[r1, c1]
    return (1 - dr) * (1 - dc) * e00 + (1 - dr) * dc * e01 + dr * (1 - dc) * e10 + dr * dc * e11

def haversine_distance(lon1, lat1, lon2, lat2):
    """计算两经纬度之间的地面大圆距离（米）"""
    R = 6371000.0  # 地球平均半径（米）
    phi1 = math.radians(lat1)
    phi2 = math.radians(lat2)
    delta_phi = math.radians(lat2 - lat1)
    delta_lambda = math.radians(lon2 - lon1)
    a = math.sin(delta_phi / 2.0)**2 + math.cos(phi1) * math.cos(phi2) * math.sin(delta_lambda / 2.0)**2
    c = 2.0 * math.atan2(math.sqrt(a), math.sqrt(1.0 - a))
    return R * c

# 2. 节点坐标与海拔定义
NODES = {
    'O01': {'name': '凤丹村无人机调度中心', 'lon': 109.2308517, 'lat': 23.0085095, 'elev': 127.7, 'type': 'hub'},
    'S001': {'name': '平旺、那龙、镇龙乡', 'lon': 109.2432319, 'lat': 23.0335927, 'elev': 154.0, 'pop': 2100},
    'S002': {'name': '上良', 'lon': 109.2585519, 'lat': 23.0703271, 'elev': 189.5, 'pop': 330},
    'S003': {'name': '六造、那旭、都长', 'lon': 109.1688578, 'lat': 23.0402714, 'elev': 308.5, 'pop': 275},
    'S004': {'name': '大站', 'lon': 109.2381714, 'lat': 23.0778414, 'elev': 198.3, 'pop': 188},
    'S005': {'name': '上良、古律、那麓', 'lon': 109.2283003, 'lat': 23.0561338, 'elev': 165.0, 'pop': 126},
    'S006': {'name': '六昌、木路', 'lon': 109.1997836, 'lat': 23.0077140, 'elev': 284.8, 'pop': 121},
    'S007': {'name': '佛子、那桑、银坑', 'lon': 109.1928191, 'lat': 23.0294186, 'elev': 318.5, 'pop': 76},
    'S008': {'name': '古麻、大站', 'lon': 109.2127472, 'lat': 23.0795113, 'elev': 240.9, 'pop': 55},
    'S009': {'name': '那龙', 'lon': 109.2541424, 'lat': 23.0552989, 'elev': 200.0, 'pop': 45},
    'S010': {'name': '镇龙乡', 'lon': 109.2760168, 'lat': 23.0194009, 'elev': 321.6, 'pop': 31},
    'S011': {'name': '那从', 'lon': 109.2203356, 'lat': 23.0319231, 'elev': 182.1, 'pop': 30},
    'S012': {'name': '镇龙乡', 'lon': 109.2831357, 'lat': 23.0544640, 'elev': 249.8, 'pop': 16},
    'S013': {'name': '镇龙乡', 'lon': 109.2713839, 'lat': 23.0310882, 'elev': 321.7, 'pop': 14},
    'S014': {'name': '镇龙乡', 'lon': 109.2835966, 'lat': 23.0052097, 'elev': 321.0, 'pop': 12},
    'S015': {'name': '六造、那托、那旭', 'lon': 109.1923786, 'lat': 23.0494548, 'elev': 444.5, 'pop': 3}
}

# 3. 机型参数
DRONE_PARAMS = {
    'A': {
        'name': '中轻载标准测试多旋翼', 'm_empty': 70.0, 'w_max': 25.0, 'vol_max': 0.060,
        'v_cruise': 12.0, 'r_empty': 25000.0, 'r_full': 20000.0, 'e_avail': 4.5, 'soc_min': 0.20,
        't_prep': 300.0, 't_load_per_box': 30.0, 't_handover_base': 150.0, 't_handover_box': 30.0,
        'v_climb': 3.0, 'v_desc': 2.5, 'eta_climb': 0.72, 'eta_desc': 0.0,
        'full_charge_time': 1800.0, 'total_batteries': 6, 'drones': ['U01', 'U02', 'U03', 'U04']
    },
    'B': {
        'name': '中载标准测试多旋翼', 'm_empty': 65.0, 'w_max': 30.0, 'vol_max': 0.073,
        'v_cruise': 15.0, 'r_empty': 28000.0, 'r_full': 16000.0, 'e_avail': 4.0, 'soc_min': 0.20,
        't_prep': 300.0, 't_load_per_box': 30.0, 't_handover_base': 150.0, 't_handover_box': 30.0,
        'v_climb': 3.0, 'v_desc': 2.5, 'eta_climb': 0.72, 'eta_desc': 0.0,
        'full_charge_time': 2400.0, 'total_batteries': 4, 'drones': ['U05', 'U06']
    },
    'C': {
        'name': '重载标准测试多旋翼', 'm_empty': 69.9, 'w_max': 80.0, 'vol_max': 0.250,
        'v_cruise': 15.0, 'r_empty': 26000.0, 'r_full': 12000.0, 'e_avail': 8.0, 'soc_min': 0.20,
        't_prep': 300.0, 't_load_per_box': 30.0, 't_handover_base': 180.0, 't_handover_box': 36.0,
        'v_climb': 2.5, 'v_desc': 2.0, 'eta_climb': 0.72, 'eta_desc': 0.0,
        'full_charge_time': 3000.0, 'total_batteries': 4, 'drones': ['U07', 'U08']
    },
    'R': {
        'name': '中继标准测试多旋翼', 'm_empty': 21.0, 'm_module': 2.5, 'm_takeoff': 23.5,
        'v_cruise': 15.0, 'p_cruise': 1.15, 'e_avail': 3.2, 'soc_min': 0.20,
        't_prep': 180.0, 't_link': 30.0, 't_turnaround': 300.0,
        'v_climb': 4.0, 'v_desc': 3.0, 'eta_climb': 0.72, 'eta_desc': 0.0,
        'p_hover': 1.05, 'p_com': 0.05, 'max_hover_agl': 300.0,
        'full_charge_time': 1800.0, 'total_modules': 6, 'drones': ['R01', 'R02']
    }
}

# 4. 加载物资清单（80个货箱）
def load_cargo_boxes():
    path = os.path.join(BASE_DATA_DIR, "物资需求与配送时限.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['逐箱货箱清单']
    boxes = []
    for r in range(2, ws.max_row + 1):
        box_id = ws.cell(r, 1).value
        if not box_id:
            continue
        svc_id = ws.cell(r, 2).value
        m_type = ws.cell(r, 3).value
        weight = float(ws.cell(r, 4).value)
        vol = float(ws.cell(r, 5).value)
        is_first = (ws.cell(r, 6).value == '是')
        first_deadline = ws.cell(r, 7).value
        first_deadline = float(first_deadline) if first_deadline is not None else None
        expect_time = float(ws.cell(r, 8).value) if ws.cell(r, 8).value is not None else None
        priority = float(ws.cell(r, 9).value)
        boxes.append({
            'box_id': box_id,
            'service_id': svc_id,
            'type': m_type,
            'weight': weight,
            'volume': vol,
            'is_first_batch': is_first,
            'first_deadline': first_deadline,
            'expect_time': expect_time,
            'priority': priority
        })
    return boxes

CARGO_BOXES = load_cargo_boxes()

# 5. 航段剖面与飞行参数缓存计算
PATH_CACHE = {}

def get_path_profile(start_node, end_node):
    """计算航段两点间直线的最高地面海拔、固定巡航海拔、起终作业高度与爬升下降高度"""
    key = (start_node, end_node)
    if key in PATH_CACHE:
        return PATH_CACHE[key]
    
    n1 = NODES[start_node]
    n2 = NODES[end_node]
    dist_m = haversine_distance(n1['lon'], n1['lat'], n2['lon'], n2['lat'])
    
    # 离散步进采样检测最高地面海拔
    num_steps = max(20, int(dist_m / 30.0))
    lons = np.linspace(n1['lon'], n2['lon'], num_steps)
    lats = np.linspace(n1['lat'], n2['lat'], num_steps)
    elevs = [get_dem_elevation(lo, la) for lo, la in zip(lons, lats)]
    max_ground_elev = max(elevs)
    
    cruise_elev = max_ground_elev + 50.0  # 计划巡航海拔
    
    # 起点终点作业高度
    z_start = n1['elev'] if start_node == 'O01' else n1['elev'] + 30.0
    z_end = n2['elev'] if end_node == 'O01' else n2['elev'] + 30.0
    
    climb_h = max(0.0, cruise_elev - z_start)
    desc_h = max(0.0, cruise_elev - z_end)
    
    res = {
        'dist_m': dist_m,
        'max_ground_elev': max_ground_elev,
        'cruise_elev': cruise_elev,
        'z_start': z_start,
        'z_end': z_end,
        'climb_h': climb_h,
        'desc_h': desc_h
    }
    PATH_CACHE[key] = res
    return res

def compute_leg_flight(m_type, start_node, end_node, payload_kg):
    """计算单航段在机型 m_type、载荷 payload_kg 下的耗时（秒）与消耗能量（kWh）"""
    prof = get_path_profile(start_node, end_node)
    params = DRONE_PARAMS[m_type]
    
    # 航段飞行耗时
    t_cruise = prof['dist_m'] / params['v_cruise']
    t_climb = prof['climb_h'] / params['v_climb']
    t_desc = prof['desc_h'] / params['v_desc']
    total_time = t_cruise + t_climb + t_desc
    
    # 等效航程与水平巡航能耗（题面官方 3/2 次方非线性动力学模型）
    w_ratio = payload_kg / params['w_max']
    r_w = params['r_empty'] - (w_ratio ** 1.5) * (params['r_empty'] - params['r_full'])
    e_cruise = (prof['dist_m'] / r_w) * params['e_avail']
    
    # 爬升附加能耗
    m_total = params['m_empty'] + payload_kg
    e_climb = (m_total * 9.80665 * prof['climb_h']) / (params['eta_climb'] * 3.6e6)
    
    total_energy = e_cruise + e_climb
    return total_time, total_energy

# 6. 两阶段电池等效充电模型
def compute_recharge_time(m_type, remaining_soc):
    """输入剩余 SOC (0.2~1.0)，输出充至 100% 所需的时间（秒）"""
    t_full = DRONE_PARAMS[m_type]['full_charge_time']
    if remaining_soc >= 1.0:
        return 0.0
    if remaining_soc < 0.90:
        frac = ((0.90 - remaining_soc) / 0.90) * 0.65 + 0.35
    else:
        frac = ((1.0 - remaining_soc) / 0.10) * 0.35
    return frac * t_full

# 7. 三维视距（LOS）空间通视判定
def check_line_of_sight(pos1, pos2):
    """
    判定三维空间点 pos1=(lon, lat, elev) 与 pos2=(lon, lat, elev) 之间是否存在山体遮挡。
    返回 True 表示无遮挡（视距通视），False 表示有山体遮挡。
    """
    lon1, lat1, z1 = pos1
    lon2, lat2, z2 = pos2
    horiz_dist = haversine_distance(lon1, lat1, lon2, lat2)
    steps = max(20, int(horiz_dist / 20.0))  # 每 20 米一个采样步长
    alphas = np.linspace(0.0, 1.0, steps)
    for a in alphas[1:-1]:  # 排除两端点
        cur_lon = (1.0 - a) * lon1 + a * lon2
        cur_lat = (1.0 - a) * lat1 + a * lat2
        cur_z = (1.0 - a) * z1 + a * z2
        dem_z = get_dem_elevation(cur_lon, cur_lat)
        if cur_z <= dem_z:
            return False  # 视线被山体阻挡
    return True

# 8. 无线通信链路损耗与可用性判定
COM_PARAMS = {
    'f_mhz': 2400.0,
    'L_sys': 3.0,
    'L_obs': 10.0,
    'P_sens': -98.0,
    'M': 8.0,
    'G01': {'Pt': 27.0, 'G': 12.0, 'hG': 20.0},
    'UAV': {'Pt': 20.0, 'G': 3.0},
    'Relay_access': {'Pt': 20.0, 'G': 6.0},
    'Relay_backhaul': {'Pt': 19.0, 'G': 8.0}
}

def free_space_path_loss(dist_km):
    if dist_km <= 0.001:
        return 0.0
    return 32.45 + 20.0 * math.log10(COM_PARAMS['f_mhz']) + 20.0 * math.log10(dist_km)

def is_link_available(pos_tx, tx_type, pos_rx, rx_type):
    """
    计算两端点之间的双向无线链路可用性
    pos=(lon, lat, alt_m)
    tx_type/rx_type: 'G01', 'UAV', 'Relay_access', 'Relay_backhaul'
    """
    lon1, lat1, z1 = pos_tx
    lon2, lat2, z2 = pos2 = pos_rx
    d_h = haversine_distance(lon1, lat1, lon2, lat2)
    d_3d_km = math.sqrt(d_h**2 + (z1 - z2)**2) / 1000.0
    
    los = check_line_of_sight(pos_tx, pos_rx)
    loss = free_space_path_loss(d_3d_km) + (0.0 if los else COM_PARAMS['L_obs'])
    
    # 判定通信方向 1 -> 2
    pt1 = COM_PARAMS[tx_type]['Pt']
    gt1 = COM_PARAMS[tx_type]['G']
    gr2 = COM_PARAMS[rx_type]['G']
    L_max_12 = pt1 + gt1 + gr2 - COM_PARAMS['P_sens'] - COM_PARAMS['M'] - COM_PARAMS['L_sys']
    
    # 判定通信方向 2 -> 1
    pt2 = COM_PARAMS[rx_type]['Pt']
    gt2 = COM_PARAMS[rx_type]['G']
    gr1 = COM_PARAMS[tx_type]['G']
    L_max_21 = pt2 + gt2 + gr1 - COM_PARAMS['P_sens'] - COM_PARAMS['M'] - COM_PARAMS['L_sys']
    
    L_max = min(L_max_12, L_max_21)
    return loss <= L_max, loss, L_max
