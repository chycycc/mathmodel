# -*- coding: utf-8 -*-
"""
全题目最终成果一键打包导出程序
读取 Q1、Q2、Q3、Q4 最新经过独立审计的官方 CSV 结果，
严格按照国赛组委会《结果提交模板.xlsx》格式规范，
生成包含全部 6 个工作表的最终交付文件：results/结果提交_最终完整版.xlsx
"""
import os
import sys
import openpyxl
import pandas as pd

_CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
_WORKSPACE_ROOT = os.path.abspath(os.path.join(_CURRENT_DIR, '..', '..'))
sys.path.insert(0, _WORKSPACE_ROOT)

sys.stdout.reconfigure(encoding='utf-8')

def export_full_excel():
    template_path = os.path.join(_WORKSPACE_ROOT, '数模题目', 'D题', '结果提交模板.xlsx')
    output_path = os.path.join(_WORKSPACE_ROOT, 'results', '结果提交_最终完整版.xlsx')
    
    wb = openpyxl.load_workbook(template_path)
    print(f"模板载入成功！工作表清单: {wb.sheetnames}")
    
    # 1. Sheet: Q1_单点组批
    ws_q1 = wb['Q1_单点组批']
    p_q1 = os.path.join(_WORKSPACE_ROOT, 'results', 'Q1', 'Q1_单点组批方案.csv')
    if not os.path.exists(p_q1):
        p_q1 = os.path.join(_WORKSPACE_ROOT, 'results', 'Q1_单点组批方案.csv')
    df_q1 = pd.read_csv(p_q1)
    # 清空第 2 行以下
    while ws_q1.max_row > 1:
        ws_q1.delete_rows(2)
    for row in df_q1.itertuples(index=False):
        ws_q1.append(list(row))
    print(f"1. Q1_单点组批 写入完成: {len(df_q1)} 行")
    
    # 2. Sheet: Q2_运输架次
    ws_q2_t = wb['Q2_运输架次']
    p_q2_t = os.path.join(_WORKSPACE_ROOT, 'results', 'Q2', 'Q2_运输架次.csv')
    if not os.path.exists(p_q2_t):
        p_q2_t = os.path.join(_WORKSPACE_ROOT, 'results', 'Q2_运输架次.csv')
    df_q2_t = pd.read_csv(p_q2_t)
    while ws_q2_t.max_row > 1:
        ws_q2_t.delete_rows(2)
    for row in df_q2_t.itertuples(index=False):
        ws_q2_t.append(list(row))
    print(f"2. Q2_运输架次 写入完成: {len(df_q2_t)} 行")
    
    # 3. Sheet: Q2_逐箱交付
    ws_q2_b = wb['Q2_逐箱交付']
    p_q2_b = os.path.join(_WORKSPACE_ROOT, 'results', 'Q2', 'Q2_逐箱交付.csv')
    if not os.path.exists(p_q2_b):
        p_q2_b = os.path.join(_WORKSPACE_ROOT, 'results', 'Q2_逐箱交付.csv')
    df_q2_b = pd.read_csv(p_q2_b)
    while ws_q2_b.max_row > 1:
        ws_q2_b.delete_rows(2)
    for row in df_q2_b.itertuples(index=False):
        ws_q2_b.append(list(row))
    print(f"3. Q2_逐箱交付 写入完成: {len(df_q2_b)} 行")
    
    # 4. Sheet: Q3_中继架次
    ws_q3_r = wb['Q3_中继架次']
    df_q3_r = pd.read_csv(os.path.join(_WORKSPACE_ROOT, 'results', 'Q3', 'Q3_中继架次.csv'))
    while ws_q3_r.max_row > 1:
        ws_q3_r.delete_rows(2)
    for row in df_q3_r.itertuples(index=False):
        ws_q3_r.append(list(row))
    print(f"4. Q3_中继架次 写入完成: {len(df_q3_r)} 行")
    
    # 5. Sheet: Q3_通信保障
    ws_q3_c = wb['Q3_通信保障']
    df_q3_c = pd.read_csv(os.path.join(_WORKSPACE_ROOT, 'results', 'Q3', 'Q3_通信保障.csv'))
    while ws_q3_c.max_row > 1:
        ws_q3_c.delete_rows(2)
    for row in df_q3_c.itertuples(index=False):
        ws_q3_c.append(list(row))
    print(f"5. Q3_通信保障 写入完成: {len(df_q3_c)} 行")
    
    # 6. Sheet: Q4_分区配置
    ws_q4 = wb['Q4_分区配置']
    df_q4 = pd.read_csv(os.path.join(_WORKSPACE_ROOT, 'results', 'Q4', 'Q4_分区配置.csv'))
    while ws_q4.max_row > 1:
        ws_q4.delete_rows(2)
    for row in df_q4.itertuples(index=False):
        ws_q4.append(list(row))
    print(f"6. Q4_分区配置 写入完成: {len(df_q4)} 行")
    
    wb.save(output_path)
    print(f"\n>>> 全部 6 个工作表已成功完整保存至: {output_path} <<<")

if __name__ == '__main__':
    export_full_excel()
