# -*- coding: utf-8 -*-
"""把问题一、问题二的三个正式 CSV 写入一个 Q1/Q2 提交工作簿。"""
from pathlib import Path
import sys
import openpyxl
import pandas as pd

RESULTS = Path(r"D:\数学建模代码\results")
TEMPLATE = Path(r"D:\数学建模代码\数模题目\D题\结果提交模板.xlsx")
OUTPUT = RESULTS / "结果提交_问题一二.xlsx"


def clear_rows(ws):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


def write_df(ws, df, columns, numeric_columns=()):
    clear_rows(ws)
    for row_idx, row in enumerate(df[columns].itertuples(index=False, name=None), start=2):
        for col_idx, (name, value) in enumerate(zip(columns, row), start=1):
            if name in numeric_columns and not pd.isna(value):
                value = float(value)
            ws.cell(row=row_idx, column=col_idx, value=value)


def main():
    required = [
        RESULTS / "Q1_单点组批方案.csv",
        RESULTS / "Q2_运输架次.csv",
        RESULTS / "Q2_逐箱交付.csv",
    ]
    missing = [str(p) for p in required if not p.exists()]
    if missing:
        raise FileNotFoundError("缺少正式结果文件：" + "; ".join(missing))

    wb = openpyxl.load_workbook(TEMPLATE)
    for sheet in ("Q3_中继架次", "Q3_通信保障", "Q4_分区配置"):
        if sheet in wb.sheetnames:
            del wb[sheet]

    q1_cols = ['架次编号', '服务区编号', '机型编号', '货箱编号列表', '总质量（kg）', '总体积（m³）', '往返时间（s）', '架次能耗（kWh）', '返航SOC（%）']
    q2t_cols = ['架次编号', '无人机编号', '机型编号', '电池编号', '开始时刻（s）', '访问服务区顺序', '返回O01时刻（s）', '架次能耗（kWh）']
    q2b_cols = ['货箱编号', '架次编号', '服务区编号', '交付完成时刻（s）']
    write_df(wb['Q1_单点组批'], pd.read_csv(required[0], encoding='utf-8-sig'), q1_cols,
             {'总质量（kg）', '总体积（m³）', '往返时间（s）', '架次能耗（kWh）', '返航SOC（%）'})
    write_df(wb['Q2_运输架次'], pd.read_csv(required[1], encoding='utf-8-sig'), q2t_cols,
             {'开始时刻（s）', '返回O01时刻（s）', '架次能耗（kWh）'})
    write_df(wb['Q2_逐箱交付'], pd.read_csv(required[2], encoding='utf-8-sig'), q2b_cols,
             {'交付完成时刻（s）'})
    wb.save(OUTPUT)
    print(f"已写入：{OUTPUT}")
    print(f"Q1={len(pd.read_csv(required[0]))} 行，Q2运输={len(pd.read_csv(required[1]))} 行，Q2逐箱={len(pd.read_csv(required[2]))} 行")


if __name__ == '__main__':
    main()
